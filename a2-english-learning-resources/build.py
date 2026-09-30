"""Build the Destination A2 practice site into ./html.
Structure follows the book's contents: 42 units, a review after every 3 units, 2 progress tests.
Extra pages: index dashboard, placement test, mistakes notebook, classroom quiz, printable worksheets,
plus Kahoot / Quizizz import spreadsheets. static/app.css + static/app.js are inlined into every page
so each HTML file also works on its own (e.g. sent via Telegram).
Run: py build.py"""
import hashlib, html, io, json, random, re, pathlib, zipfile
from units_a import UNITS_A
from units_b import UNITS_B
from units_c import UNITS_C
from readings_a import READINGS_A
from readings_b import READINGS_B
from review_items import REVIEWS, PROGRESS

ROOT = pathlib.Path(__file__).parent
OUT = ROOT / "html"
CSS = (ROOT / "static" / "app.css").read_text(encoding="utf-8")
JS = (ROOT / "static" / "app.js").read_text(encoding="utf-8")
UNITS = sorted(UNITS_A + UNITS_B + UNITS_C, key=lambda u: u["n"])
BY_N = {u["n"]: u for u in UNITS}
READINGS = {**READINGS_A, **READINGS_B}
NO_ARTICLE = {"-", "—", "no article", "0", "x"}
EMOJI = dict(enumerate("🏠 🎬 👪 🕰️ 📜 🎨 🌧️ 📣 ⚽ ✅ ⏳ 🏫 🔮 🙋 💼 🗓️ 🚦 🛒 🏗️ 🔤 🎉 🍞 ⚖️ 👗 📏 🏆 🏙️ 🏃 🧭 ✈️ 🔀 💭 💪 🤝 ❓ 💻 🔗 ⏮️ 🌳 💬 🔄 🌍".split(), 1))
esc = lambda s: html.escape(s, quote=True)
hue = lambda n: (n * 37 + 200) % 360
kind_of = lambda u: "grammar" if "lesson" in u else "vocab"

# Every rendered question gets a stable id so the mistakes notebook can find it again.
REG = {}                       # id -> {q, a, t, u (label), h (href)}
CUR = {"label": "", "href": ""}


def iid(scope, q, a):
    return "i" + hashlib.md5(f"{scope}|{q}|{a[0]}".encode("utf-8")).hexdigest()[:10]


def register(scope, q, a, t):
    k = iid(scope, q, a)
    REG.setdefault(k, {"q": q, "a": a, "t": t, "u": CUR["label"], "h": CUR["href"]})
    return k


def fill(q, a):
    """Turn a gap / multiple-choice item into a full sentence for the games."""
    s = q.replace("___ ", "").replace("___", "") if a in NO_ARTICLE else q.replace("___", a, 1)
    s = re.sub(r"\s*\([^)]*\)", "", s)                     # drop hints like (not like)
    s = re.sub(r" '(ll|m|re|s|ve|d)\b", r"'\1", s)          # I 'll -> I'll
    return re.sub(r"\s+", " ", s).strip()


def sentences(u):
    items = [(q, a[0]) for ex in u["ex"] if ex[1] == "mc" for q, a in ex[2]] + [(q, a[0]) for q, a in u["test"]]
    items += [(q, next((x for x in a if not x.startswith("'")), a[0])) for ex in u["ex"] if ex[1] == "gap" for q, a in ex[2]]
    out = []
    for q, a in items:
        s = fill(q, a)
        if "→" in s or '"' in s or "___" in s or s[-1:] not in ".?!" or len(s.split()) < 3 or s in out:
            continue
        out.append(s)
    return out


def mc_pool(u):
    return [it for ex in u["ex"] if ex[1] == "mc" for it in ex[2]] + u["test"]


def validate():
    assert [u["n"] for u in UNITS] == list(range(1, 43)), "units must be numbered 1..42"
    unit_qs = set()
    for u in UNITS:
        assert len(u["ex"]) == 3, f"unit {u['n']}: needs exercises A, B, C"
        assert len(u["test"]) == 10, f"unit {u['n']}: test needs 10 questions"
        for ex in u["ex"]:
            for q, a in ex[2]:
                unit_qs.add(q)
                if ex[1] == "mc":
                    assert len(a) == 3 and len(set(a)) == 3, f"unit {u['n']}: bad options {q}"
                else:
                    assert q.count("___") == 1 and a and all(a), f"unit {u['n']}: bad gap {q}"
        for q, a in u["test"]:
            unit_qs.add(q)
            assert len(a) == 3 and len(set(a)) == 3, f"unit {u['n']}: bad test options {q}"
        assert len(sentences(u)) >= 8, f"unit {u['n']}: too few sentences for the games"
    assert sorted(READINGS) == list(range(1, 43)), "every unit needs a reading"
    for n, r in READINGS.items():
        assert r["title"] and len(r["text"]) > 300 and len(r["questions"]) == 5, f"reading {n}"
        for q, a in r["questions"]:
            assert len(a) == 3 and len(set(a)) == 3, f"reading {n}: bad options {q}"
    assert sorted(REVIEWS) == list(range(1, 15)) and sorted(PROGRESS) == [1, 2]
    new_qs = []
    for k, r in REVIEWS.items():
        assert len(r["mc"]) == 12 and len(r["gap"]) == 6, f"review {k}: needs 12 mc + 6 gap"
        for q, a in r["mc"]:
            assert len(a) == 3 and len(set(a)) == 3, f"review {k}: bad options {q}"
        for q, a in r["gap"]:
            assert q.count("___") == 1 and a and all(a), f"review {k}: bad gap {q}"
        new_qs += [q for q, _ in r["mc"] + r["gap"]]
    for k, items in PROGRESS.items():
        assert len(items) == 30, f"progress test {k}: needs 30 questions"
        for q, a in items:
            assert len(a) == 3 and len(set(a)) == 3, f"progress {k}: bad options {q}"
        new_qs += [q for q, _ in items]
    dup = [q for q in new_qs if q in unit_qs] + [q for q in set(new_qs) if new_qs.count(q) > 1]
    assert not dup, f"review/progress questions must be new: {dup}"


def shuffled(block, i, opts):
    s = opts[:]
    random.Random(f"{block}-{i}").shuffle(s)
    return s


def blank(q, mark='<span class="blank">______</span>'):
    return esc(q).replace("___", mark)


def mc(block, items, scope, units=None):
    rows = []
    for i, (q, opts) in enumerate(items):
        k = register(scope if isinstance(scope, str) else scope[i], q, opts, "mc")
        labels = "".join(f'<label><input type="radio" name="{block}-{i}" value="{int(o == opts[0])}"> {esc(o)}</label>'
                         for o in shuffled(block, i, opts))
        du = f' data-u="{units[i]}"' if units else ""
        rows.append(f'<li class="q" data-id="{k}"{du}><p>{blank(q)}</p><div class="opts">{labels}</div></li>')
    return "<ol>" + "".join(rows) + "</ol>"


def gap(items, scope):
    rows = []
    for i, (q, answers) in enumerate(items):
        k = register(scope if isinstance(scope, str) else scope[i], q, answers, "gap")
        box = f'<input type="text" autocomplete="off" spellcheck="false" data-a="{esc(json.dumps(answers))}">'
        rows.append(f'<li class="q" data-id="{k}"><p>{esc(q).replace("___", box)}</p></li>')
    return "<ol>" + "".join(rows) + "</ol>"


BUTTONS = ('<div class="actions"><button onclick="check(this)">Check</button>'
           '<button class="ghost" onclick="reveal(this)">Show answers</button>'
           '<button class="ghost" onclick="reset(this)">Try again</button><span class="score"></span></div>')
FINISH = ('<div class="actions"><button onclick="finish(this)">Finish</button>'
          '<button class="ghost hidden" onclick="reveal(this)">Show answers</button>'
          '<button class="ghost" onclick="reset(this)">Try again</button><span class="score"></span></div>')


def lesson_html(u, speak=True):
    if "lesson" in u:
        parts = []
        for head, lines in u["lesson"]:
            lis = "".join(f'<li class="uz">{l[4:]}</li>' if l.startswith("UZ: ") else f"<li>{l}</li>" for l in lines)
            parts.append(f"<h3>{esc(head)}</h3><ul>{lis}</ul>")
        return "".join(parts)
    sb = (lambda t: f'<button class="say" data-say="{esc(t)}" title="Listen">🔊</button>') if speak else (lambda t: "")
    rows = "".join(f"<tr><td>{sb(en)}<b>{esc(en)}</b></td><td class='uzc'>{esc(uz)}</td><td>{sb(ex)}<i>{esc(ex)}</i></td></tr>"
                   for en, uz, ex in u["words"])
    return (f"<table><thead><tr><th>Word</th><th>O'zbekcha</th><th>Example</th></tr></thead><tbody>{rows}</tbody></table>"
            f'<p class="tip"><b>Tip:</b> {esc(u["tip"])}</p>')


def paras_html(text):
    return "".join(f"<p>{esc(p).replace(chr(10), '<br>')}</p>" for p in text.split("\n\n"))


def reading_html(n):
    r = READINGS[n]
    return (f'<section id="reading" class="card"><h2>📰 {esc(r["title"])}</h2>'
            '<div class="actions"><button onclick="readAloud(this)">🔊 Listen</button>'
            '<button class="ghost" onclick="readAloud(this, 0.7)">🐢 Slowly</button>'
            '<button class="ghost" onclick="speechSynthesis.cancel()">⏹ Stop</button>'
            '<button class="ghost" onclick="toggleText(this)">🎧 Listening mode</button></div>'
            f'<div class="rtext">{paras_html(r["text"])}</div></section>'
            '<section id="readq" class="card"><h2>Comprehension</h2>'
            '<p class="inst">Listen or read, then choose the correct answer.</p>'
            f'{mc(f"rd{n}", r["questions"], f"rd{n}")}{FINISH}</section>')


def unit_file(n):
    return f"unit-{n:02d}.html"


def build_sequence():
    """(file, label, kind, unit numbers) in book order: units, Review after every 3rd, progress tests after 21 and 42."""
    seq = []
    for u in UNITS:
        n = u["n"]
        seq.append((unit_file(n), f"Unit {n}", "unit", [n]))
        if n % 3 == 0:
            seq.append((f"review-{n // 3:02d}.html", f"Review {n // 3}", "review", [n - 2, n - 1, n]))
        if n in (21, 42):
            k = 1 if n == 21 else 2
            seq.append((f"progress-test-{k}.html", f"Progress Test {k}", "progress", list(range(n - 20, n + 1))))
    return seq


SEQ = build_sequence()


def pager(fname):
    i = [s[0] for s in SEQ].index(fname)
    prev_link = f'<a href="{SEQ[i - 1][0]}">← {SEQ[i - 1][1]}</a>' if i > 0 else "<span></span>"
    next_link = f'<a href="{SEQ[i + 1][0]}">{SEQ[i + 1][1]} →</a>' if i + 1 < len(SEQ) else "<span></span>"
    return prev_link, next_link


def doc(title, body, page, h, extra_script="", body_class=""):
    return (f'<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
            f'<title>{title} — Destination A2 practice</title><style>{CSS}</style></head>'
            f'<body data-page="{page}" class="{body_class}" style="--h:{h}"><div class="wrap">{body}</div>'
            f'{extra_script}<script>{JS}</script></body></html>')


def data_script(name, data):
    return f"<script>const {name}=" + json.dumps(data, ensure_ascii=False).replace("</", "<\\/") + ";</script>"


SOUND_BTN = '<button class="pill sound" id="sound" onclick="toggleSound(this)" title="Sound on/off">🔊</button>'


def unit_page(u):
    n, kind = u["n"], kind_of(u)
    scope = f"u{n}"
    sections = []
    for k, ex in enumerate(u["ex"]):
        title, typ, items = ex[:3]
        bank = f'<p class="bank">{" · ".join(esc(w) for w in ex[3])}</p>' if len(ex) > 3 else ""
        body = mc(f"u{n}e{k}", items, scope) if typ == "mc" else gap(items, scope)
        sections.append(f'<section id="ex{k + 1}" class="card"><h2>Exercise {"ABC"[k]}</h2>'
                        f'<p class="inst">{esc(title)}</p>{bank}{body}{BUTTONS}</section>')
    test = (f'<section id="test" class="card test"><h2>📝 Unit {n} Test</h2>'
            f'<p class="inst">Choose the correct answer. 10 questions — 1 point each. 90% = ★★★</p>{mc(f"u{n}t", u["test"], scope)}{FINISH}</section>')
    prev_link, next_link = pager(unit_file(n))
    data = {"kind": kind, "words": u.get("words", []), "sentences": sentences(u),
            "mc": [[q, a, iid(scope, q, a)] for q, a in mc_pool(u)]}
    body = f"""
<header class="hero"><a class="home" href="index.html">← All units</a>{SOUND_BTN}
<div class="emoji">{EMOJI[n]}</div><div class="tag">Unit {n} · {"Grammar" if kind == "grammar" else "Vocabulary"}</div>
<h1>{esc(u["title"])}</h1>
<div class="hero-row"><span class="stars" id="hero-stars">☆☆☆</span><span class="pill" id="hero-xp">⚡ 0 XP</span>
<a class="pill" href="worksheet-{n:02d}.html">🖨 Worksheet</a><a class="pill" href="quiz.html?units={n}">🎮 Class quiz</a></div></header>
<nav class="tabs"><button data-go="lesson">📖 Lesson</button><button data-go="practice">✏️ Practice</button>
<button data-go="read">📰 Read &amp; listen</button><button data-go="games">🎮 Games</button>
<button data-go="listen">🎧 Listen &amp; speak</button><button data-go="test">📝 Test</button><button data-go="task">💬 Task</button></nav>
<section id="lesson" class="card" data-tab="lesson"><h2>📖 Lesson</h2>{lesson_html(u)}</section>
<div data-tab="practice">{"".join(sections)}</div>
<div data-tab="read">{reading_html(n)}</div>
<div id="games-area" data-tab="games"></div>
<div id="listen-area" data-tab="listen"></div>
<div data-tab="test">{test}</div>
<section id="task" class="card task" data-tab="task"><h2>💬 Speaking &amp; writing</h2><p>{esc(u["task"])}</p></section>
<div class="pager">{prev_link}{next_link}</div>
<footer>Original practice material for A2 learners, organised by the units of the book. Not copied from the coursebook.</footer>"""
    return doc(f"Unit {n}: {esc(u['title'])}", body, unit_file(n)[:-5], hue(n), data_script("UNIT", data))


def practice_page(fname, label, kind, nums):
    """Reviews and progress tests use their own new questions (review_items.py)."""
    if kind == "review":
        k = nums[-1] // 3
        mc_items, gap_items = REVIEWS[k]["mc"], REVIEWS[k]["gap"]
        scope = f"r{k}"
        heading, emoji = f"{label}: Units {nums[0]}, {nums[1]} and {nums[2]}", "🔁"
    else:
        k = 1 if nums[-1] == 21 else 2
        mc_items, gap_items = PROGRESS[k], []
        scope = f"p{k}"
        heading, emoji = f"{label}: Units {nums[0]}–{nums[-1]}", "🏁"
    sections = (f'<section id="p1" class="card test"><h2>Part 1</h2><p class="inst">Choose the correct answer. '
                f'{len(mc_items)} questions — 1 point each.</p>{mc(fname[:-5], mc_items, scope)}{FINISH}</section>')
    if gap_items:
        sections += (f'<section id="p2" class="card"><h2>Part 2</h2><p class="inst">Complete the sentences.</p>'
                     f'{gap(gap_items, scope)}{BUTTONS}</section>')
    units_list = " · ".join(f'<a href="{unit_file(n)}">{n}. {esc(BY_N[n]["title"])}</a>' for n in nums)
    prev_link, next_link = pager(fname)
    body = f"""
<header class="hero"><a class="home" href="index.html">← All units</a>{SOUND_BTN}
<div class="emoji">{emoji}</div><div class="tag">Revision · new questions</div><h1>{esc(heading)}</h1>
<div class="hero-row"><span class="stars" id="hero-stars">☆☆☆</span><span class="pill" id="hero-xp">⚡ 0 XP</span>
<a class="pill" href="quiz.html?units={",".join(map(str, nums))}">🎮 Class quiz</a></div></header>
<section class="card"><p class="inst">Units in this {"review" if kind == "review" else "test"}:</p><p>{units_list}</p></section>
{sections}
<div class="pager">{prev_link}{next_link}</div>
<footer>New revision questions for these units. Original material, not copied from the coursebook.</footer>"""
    return doc(esc(heading), body, fname[:-5], 150 if kind == "review" else 30)


def placement_page():
    """One unit-test question per unit (42); data-u lets the page map results back to units."""
    rng = random.Random("placement")
    items = [rng.choice(u["test"]) for u in UNITS]
    units = [u["n"] for u in UNITS]
    titles = {u["n"]: u["title"] for u in UNITS}
    body = f"""
<header class="hero"><a class="home" href="index.html">← All units</a>{SOUND_BTN}
<div class="emoji">🧭</div><div class="tag">Start here · Shu yerdan boshlang</div><h1>Placement test</h1>
<div class="hero-row"><span class="pill">42 questions · about 15 minutes</span><span class="pill">Find out which units to study</span></div></header>
<section class="card"><p>Answer every question without help. At the end you get a study plan: which units you already know and which ones to study first.</p>
<p class="tip">Har bir savolga yordamsiz javob bering. Oxirida qaysi unitlarni bilishingiz va qaysilarini o'qish kerakligi ko'rsatiladi.</p></section>
<section id="placement" class="card test"><h2>Placement test</h2>{mc("pl", items, [f"u{n}" for n in units], units)}
<div class="actions"><button onclick="placement(this)">See my plan</button><span class="score"></span></div></section>
<div id="plan"></div>"""
    return doc("Placement test", body, "placement", 280, data_script("TITLES", titles))


def mistakes_page():
    body = f"""
<header class="hero"><a class="home" href="index.html">← All units</a>{SOUND_BTN}
<div class="emoji">📒</div><div class="tag">Xatolar daftari</div><h1>My mistakes</h1>
<div class="hero-row"><span class="pill">Every wrong answer is saved here</span><span class="pill">Get it right on 3 different days to master it</span></div></header>
<div id="mistakes"></div>"""
    return doc("My mistakes", body, "mistakes", 20, data_script("BANK", REG))


def mcq_sheet(block, items):
    """Printable multiple-choice list + answer letters, same option order as the web page."""
    line = '<span class="line"></span>'
    rows, ans = [], []
    for i, (q, a) in enumerate(items):
        sh = shuffled(block, i, a)
        rows.append(f'<li><p class="mcq">{blank(q, line)}</p><p class="mco">'
                    + " &nbsp; ".join(f"{'abc'[j]}) {esc(o)}" for j, o in enumerate(sh)) + "</p></li>")
        ans.append(f"{i + 1} {'abc'[sh.index(a[0])]}")
    return f"<ol>{''.join(rows)}</ol>", ", ".join(ans)


def worksheet_page(u):
    n = u["n"]
    line = '<span class="line"></span>'
    parts, key = [], []
    for k, ex in enumerate(u["ex"]):
        title, typ, items = ex[:3]
        letter = "ABC"[k]
        bank = f'<p class="bank">{" · ".join(esc(w) for w in ex[3])}</p>' if len(ex) > 3 else ""
        if typ == "mc":
            body, ans = mcq_sheet(f"u{n}e{k}", items)
        else:
            body = "<ol>" + "".join(f"<li>{blank(q, line)}</li>" for q, a in items) + "</ol>"
            ans = ", ".join(f"{i + 1} {esc(a[0])}" for i, (q, a) in enumerate(items))
        parts.append(f"<h2>Exercise {letter}. {esc(title)}</h2>{bank}{body}")
        key.append(f"<p><b>Exercise {letter}:</b> {ans}</p>")
    r = READINGS[n]
    rbody, rans = mcq_sheet(f"rd{n}", r["questions"])
    tbody, tans = mcq_sheet(f"u{n}t", u["test"])
    key += [f"<p><b>Reading:</b> {rans}</p>", f"<p><b>Test:</b> {tans}</p>"]
    body = f"""
<div class="noprint actions"><button onclick="print()">🖨 Print</button><a href="{unit_file(n)}">← Back to Unit {n}</a></div>
<h1>Destination A2 · Unit {n}: {esc(u["title"])}</h1>
<div class="who"><span>Name</span><span>Class</span><span>Date</span><span>Score: ____ / {sum(len(e[2]) for e in u["ex"]) + 15}</span></div>
<h2>Lesson</h2>{lesson_html(u, speak=False)}
{"".join(parts)}
<h2>Reading: {esc(r["title"])}</h2>{paras_html(r["text"])}{rbody}
<h2>Test</h2>{tbody}
<h2>Speaking &amp; writing</h2><p>{esc(u["task"])}</p>{'<div class="wl"></div>' * 8}
<div class="page-break"></div><h1>Answer key · Unit {n}</h1><div class="key">{"".join(key)}</div>"""
    return doc(f"Worksheet Unit {n}", body, "worksheet", hue(n), body_class="sheet")


def quiz_page():
    data = [{"n": u["n"], "title": u["title"], "items": mc_pool(u)} for u in UNITS]
    body = f"""
<header class="hero"><a class="home" href="index.html">← All units</a>{SOUND_BTN}
<div class="emoji">🎮</div><div class="tag">For teachers · Sinf uchun</div><h1>Classroom quiz</h1>
<div class="hero-row"><span class="pill">Show it on the projector. Teams answer, you give points. Space = next.</span></div></header>
<div id="quiz"></div>"""
    return doc("Classroom quiz", body, "quiz", 330, data_script("QUIZ", data))


# ---------- Kahoot / Quizizz spreadsheets ----------
KAHOOT_Q, KAHOOT_A = 95, 60        # limits printed in Kahoot's own template (header row 8, data from row 9)


def kahoot_xlsx(items, title):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws["B2"] = "Quiz template"
    ws["B3"] = title
    ws["B4"] = "Made from Destination A2 practice. Upload in Kahoot: Create > Add question > Import spreadsheet."
    headers = ["Question - max 95 characters", "Answer 1 - max 60 characters", "Answer 2 - max 60 characters",
               "Answer 3 - max 60 characters", "Answer 4 - max 60 characters",
               "Time limit (sec) - 5,10,20,30,60,90 or 120 secs", "Correct answer(s) - choose at least one"]
    for c, h in enumerate(headers, start=2):
        ws.cell(row=8, column=c, value=h)
    for i, (q, opts) in enumerate(items):
        row, sh = 9 + i, shuffled(f"k-{title}", i, opts)
        ws.cell(row=row, column=1, value=i + 1)
        ws.cell(row=row, column=2, value=q.replace("___", "_____"))
        for j, o in enumerate(sh):
            ws.cell(row=row, column=3 + j, value=o)
        ws.cell(row=row, column=7, value=20)
        ws.cell(row=row, column=8, value=str(sh.index(opts[0]) + 1))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def quizizz_xlsx(items, title):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Question Text", "Question Type", "Option 1", "Option 2", "Option 3", "Option 4", "Option 5",
               "Correct Answer", "Time in seconds", "Image Link", "Answer explanation"])
    for i, (q, opts) in enumerate(items):
        sh = shuffled(f"z-{title}", i, opts)
        ws.append([q.replace("___", "_____"), "Multiple Choice", *sh, None, None, sh.index(opts[0]) + 1, 30, None, None])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_spreadsheets():
    """Writes the .xlsx / .zip files, plus xlsx-files.json (base64) because the published artifact
    cannot serve .xlsx/.zip: the online page decodes it and hands the bytes to the downloads capability."""
    import base64
    skipped = 0
    zips = {d: zipfile.ZipFile(OUT / f"{d}-all-units.zip", "w", zipfile.ZIP_DEFLATED) for d in ("kahoot", "quizizz")}
    for d in zips:
        (OUT / d).mkdir(exist_ok=True)
    for u in UNITS:
        items = mc_pool(u)
        fit = [(q, a) for q, a in items if len(q.replace("___", "_____")) <= KAHOOT_Q and max(map(len, a)) <= KAHOOT_A]
        skipped += len(items) - len(fit)
        title, name = f"Destination A2 · Unit {u['n']}: {u['title']}", f"unit-{u['n']:02d}.xlsx"
        for d, blob in (("kahoot", kahoot_xlsx(fit, title)), ("quizizz", quizizz_xlsx(items, title))):
            (OUT / d / name).write_bytes(blob)
            zips[d].writestr(name, blob)
    for z in zips.values():
        z.close()
    files = sorted(p for p in OUT.rglob("*") if p.suffix in (".xlsx", ".zip"))
    packed = {p.relative_to(OUT).as_posix(): base64.b64encode(p.read_bytes()).decode("ascii") for p in files}
    (OUT / "xlsx-files.json").write_text(json.dumps(packed) + "\n", encoding="utf-8")
    return skipped


def index_body():
    blocks, cur = [], []
    for fname, label, kind, nums in SEQ:
        if kind == "unit":
            u = BY_N[nums[0]]
            sub = "Grammar" if kind_of(u) == "grammar" else "Vocabulary"
            cur.append(f'<a class="unit" href="{fname}" data-page="{fname[:-5]}" data-kind="{kind_of(u)}" style="--h:{hue(u["n"])}">'
                       f'<span class="ic">{EMOJI[u["n"]]}</span><div><span>{label} · {sub}</span><b>{esc(u["title"])}</b>'
                       f'<span class="stars">☆☆☆</span></div></a>')
        else:
            span = f"Units {nums[0]}, {nums[1]}, {nums[2]}" if kind == "review" else f"Units {nums[0]}–{nums[-1]}"
            cur.append(f'<a class="unit r" href="{fname}" data-page="{fname[:-5]}" data-kind="{kind}" style="--h:150">'
                       f'<span class="ic">{"🔁" if kind == "review" else "🏁"}</span><div><span>{label}</span><b>{span}</b>'
                       f'<span class="stars">☆☆☆</span></div></a>')
            title = f"Units {nums[0]}–{nums[-1]}" if kind == "review" else label
            blocks.append(f'<div class="block">{title}</div><div class="grid">{"".join(cur)}</div>')
            cur = []
    sheets = " · ".join(f'<a href="worksheet-{u["n"]:02d}.html">{u["n"]}</a>' for u in UNITS)
    opts = "".join(f'<option value="{u["n"]:02d}">Unit {u["n"]}: {esc(u["title"])}</option>' for u in UNITS)
    return f"""
<header class="hero">{SOUND_BTN}<div class="emoji">🎓</div><div class="tag">A2 · Grammar &amp; Vocabulary</div>
<h1>Destination A2 — Unit Practice</h1>
<div class="hero-row"><span class="pill">42 units · 14 reviews · 2 progress tests</span><span class="pill">Lessons · reading · games · tests</span></div></header>
<div class="start">
<a class="cta" href="placement.html"><span>🧭</span><div><b>Placement test</b><small>New here? Find out where to start.</small></div></a>
<a class="cta" href="mistakes.html"><span>📒</span><div><b>My mistakes</b><small id="st-due">Your wrong answers, ready to review.</small></div></a>
</div>
<div class="dash" id="dash">
<div class="stat"><span>⚡ Experience</span><b id="st-xp">0</b><span>XP</span></div>
<div class="stat"><span>🎖 Level</span><b id="st-level">Level 1</b><span id="st-level-name"></span></div>
<div class="stat"><span>⭐ Stars</span><b id="st-stars">0</b><span>of 174</span></div>
<div class="stat"><span>🔥 Streak</span><b id="st-streak">0</b><span>days in a row</span></div>
<div class="stat"><span>✅ Unit tests</span><b id="st-units">0 / 42</b><div class="bar"><i id="bar-units"></i></div></div>
</div>
<section class="card"><h2>🏅 Badges</h2><div class="badges" id="badges"></div></section>
<section class="card teacher"><h2>👩‍🏫 For teachers · O'qituvchilar uchun</h2>
<p>🎮 <a href="quiz.html"><b>Classroom quiz</b></a> — Kahoot-style team game for the projector: choose units, teams and time.</p>
<p>🖨 <b>Printable worksheets</b> with reading and answer key — Unit: {sheets}</p>
<p>📥 <b>Kahoot / Quizizz files</b> — each unit's 16 test questions as a spreadsheet to import.</p>
<div class="actions"><select id="xl-unit" aria-label="Unit">{opts}</select>
<a class="dl btnlink" data-dir="kahoot" href="kahoot/unit-01.xlsx" download>Kahoot (.xlsx)</a>
<a class="dl btnlink" data-dir="quizizz" href="quizizz/unit-01.xlsx" download>Quizizz (.xlsx)</a></div>
<div class="actions"><a class="dl btnlink ghost" href="kahoot-all-units.zip" download>All units for Kahoot (.zip)</a>
<a class="dl btnlink ghost" href="quizizz-all-units.zip" download>All units for Quizizz (.zip)</a><span class="fb" id="dl-msg"></span></div>
<p class="inst">Kahoot: Create → Add question → Import spreadsheet. Quizizz / Wayground: Create → Assessment → Import from spreadsheet.</p></section>
{"".join(blocks)}
<footer>Progress is saved in this browser only. <button class="ghost" id="reset">Reset my progress</button><br>
Original practice material for A2 learners, organised by the units of the book. Not copied from the coursebook.</footer>"""


def index_page():
    return doc("Home", index_body(), "index", 230)


def artifact_home():
    """Entry page for the published artifact: the viewer adds <html>/<head>/<body> itself."""
    return (f'<title>Destination A2 Practice</title><style>{CSS}</style>'
            f'<div class="wrap">{index_body()}</div><script>document.body.dataset.page="index";</script><script>{JS}</script>')


def write(fname, text):
    (OUT / fname).write_text(text, encoding="utf-8")


if __name__ == "__main__":
    validate()
    OUT.mkdir(exist_ok=True)
    for fname, label, kind, nums in SEQ:
        CUR.update(label=label if kind != "unit" else f"Unit {nums[0]}: {BY_N[nums[0]]['title']}", href=fname)
        write(fname, unit_page(BY_N[nums[0]]) if kind == "unit" else practice_page(fname, label, kind, nums))
    for u in UNITS:
        write(f"worksheet-{u['n']:02d}.html", worksheet_page(u))
    CUR.update(label="Placement test", href="placement.html")
    write("placement.html", placement_page())
    write("quiz.html", quiz_page())
    write("mistakes.html", mistakes_page())          # last: embeds every registered question
    write("index.html", index_page())
    (ROOT / "artifact-home.html").write_text(artifact_home(), encoding="utf-8")
    skipped = export_spreadsheets()
    print(f"built {len(SEQ)} pages ({len(UNITS)} units), {len(UNITS)} worksheets, placement, mistakes, quiz, index; "
          f"{len(REG)} questions registered; Kahoot skipped {skipped} over-length questions -> {OUT}")
